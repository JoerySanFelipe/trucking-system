import docx

def read_docx(file_path):
    doc = docx.Document(file_path)
    lines = []
    lines.append("=== PARAGRAPHS ===")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            lines.append(f"P{i+1}: {p.text}")
            
    lines.append("\n=== TABLES ===")
    for t_idx, table in enumerate(doc.tables):
        lines.append(f"\n--- TABLE {t_idx+1} ---")
        for row in table.rows:
            row_text = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
            lines.append(" | ".join(row_text))
            
    with open(r"c:\kudecode\porbido-trucking\docs\extracted_docx_text.txt", "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print("Extracted UTF-8 text successfully saved!")

if __name__ == "__main__":
    read_docx(r"c:\kudecode\porbido-trucking\docs\Porbido_Commercial_System_Proposal.docx")
